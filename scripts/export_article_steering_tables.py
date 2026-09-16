#!/usr/bin/env python3
"""Export publication-ready single/dual atomic steering tables.

The input is the cluster-seen variant/aggregate CSV pair produced by
``analyze_cluster_seen_atomic_sweep.py``.  Percentages are written as numeric
Excel cells (not strings) so that authors can change decimal precision later.
"""

from __future__ import annotations

import argparse
import csv
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


METRICS = (
    ("pair_steer_success_rate", "Atomic↔Reverse (%)"),
    ("steer_vs_empty_success_rate", "Atomic vs Empty (%)"),
    ("prompt_absolute_success_rate", "Atomic absolute direction (%)"),
)

ATOM_SYMBOL = {
    "move_x_neg": "T_x−",
    "move_x_pos": "T_x+",
    "move_y_neg": "T_y−",
    "move_y_pos": "T_y+",
    "move_z_neg": "T_z−",
    "move_z_pos": "T_z+",
    "rotate_x_neg": "R_x−",
    "rotate_x_pos": "R_x+",
    "rotate_y_neg": "R_y−",
    "rotate_y_pos": "R_y+",
    "rotate_z_neg": "R_z−",
    "rotate_z_pos": "R_z+",
}


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser()
    parser.add_argument("--report-dir", type=Path, required=True)
    parser.add_argument("--output-dir", type=Path, required=True)
    parser.add_argument("--step", type=int, default=50)
    return parser.parse_args()


def read_csv(path: Path) -> list[dict[str, str]]:
    with path.open(newline="", encoding="utf-8") as handle:
        return list(csv.DictReader(handle))


def pct(row: dict[str, str], key: str) -> float | None:
    value = row[key].strip()
    return 100.0 * float(value) if value else None


def symbol(label: str) -> str:
    return " + ".join(ATOM_SYMBOL[item] for item in label.split("+"))


def normalized_rows(rows: list[dict[str, str]]) -> list[dict[str, object]]:
    result = []
    for row in rows:
        result.append(
            {
                "canonical": row["requested_atoms"],
                "symbol": symbol(row["requested_atoms"]),
                "trials": int(row["trials"]),
                "paired_trials": int(row["paired_trials"]),
                "reverse": pct(row, "pair_steer_success_rate"),
                "empty": pct(row, "steer_vs_empty_success_rate"),
                "absolute": pct(row, "prompt_absolute_success_rate"),
            }
        )
    return result


def write_csv(path: Path, rows: list[dict[str, object]]) -> None:
    fields = (
        "canonical",
        "symbol",
        "trials",
        "paired_trials",
        "reverse",
        "empty",
        "absolute",
    )
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields)
        writer.writeheader()
        writer.writerows(rows)


def add_sheet(workbook: Workbook, title: str, rows: list[dict[str, object]]) -> None:
    worksheet = workbook.create_sheet(title)
    headers = [
        "Atomic prompt",
        "Symbol",
        "Trials",
        "Paired trials",
        "Atomic↔Reverse (%)",
        "Atomic vs Empty (%)",
        "Atomic absolute direction (%)",
    ]
    worksheet.append(headers)
    for row in rows:
        worksheet.append(
            [
                row["canonical"],
                row["symbol"],
                row["trials"],
                row["paired_trials"],
                row["reverse"] / 100.0 if row["reverse"] is not None else None,
                row["empty"] / 100.0 if row["empty"] is not None else None,
                row["absolute"] / 100.0 if row["absolute"] is not None else None,
            ]
        )
    for cell in worksheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="355C7D")
        cell.alignment = Alignment(horizontal="center", vertical="center")
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    for column in range(1, 8):
        worksheet.column_dimensions[get_column_letter(column)].width = (
            37 if column == 1 else 29 if column == 7 else 22
        )
    for row in worksheet.iter_rows(min_row=2, min_col=5, max_col=7):
        for cell in row:
            cell.number_format = "0.00%"
            cell.alignment = Alignment(horizontal="center")


def markdown_table(rows: list[dict[str, object]]) -> str:
    lines = [
        "| Atomic prompt | Symbol | Trials | Paired trials | Atomic↔Reverse | Atomic vs Empty | Atomic absolute direction |",
        "|---|---:|---:|---:|---:|---:|---:|",
    ]
    for row in rows:
        reverse = "—" if row["reverse"] is None else f"{row['reverse']:.2f}%"
        empty = "—" if row["empty"] is None else f"{row['empty']:.2f}%"
        absolute = (
            "—" if row["absolute"] is None else f"{row['absolute']:.2f}%"
        )
        lines.append(
            f"| {row['canonical']} | {row['symbol']} | {row['trials']} | {row['paired_trials']} | "
            f"{reverse} | {empty} | {absolute} |"
        )
    return "\n".join(lines)


def latex_escape(value: str) -> str:
    return value.replace("_", r"\_")


def latex_table(title: str, label: str, rows: list[dict[str, object]]) -> str:
    body = []
    for row in rows:
        reverse = "--" if row["reverse"] is None else f"{row['reverse']:.2f}"
        empty = "--" if row["empty"] is None else f"{row['empty']:.2f}"
        absolute = "--" if row["absolute"] is None else f"{row['absolute']:.2f}"
        body.append(
            f"{latex_escape(str(row['canonical']))} & {row['trials']} & {row['paired_trials']} & "
            f"{reverse} & {empty} & {absolute} \\\\"
        )
    return "\n".join(
        [
            r"\begin{longtable}{lrrrrr}",
            f"\\caption{{{title}}}\\label{{{label}}} \\\\",
            r"\toprule",
            r"Atomic prompt & Trials & Paired & A$\leftrightarrow$R (\%) & A vs. Empty (\%) & Absolute (\%) \\",
            r"\midrule",
            r"\endfirsthead",
            r"\toprule",
            r"Atomic prompt & Trials & Paired & A$\leftrightarrow$R (\%) & A vs. Empty (\%) & Absolute (\%) \\",
            r"\midrule",
            r"\endhead",
            *body,
            r"\bottomrule",
            r"\end{longtable}",
        ]
    )


def main() -> None:
    args = parse_args()
    args.output_dir.mkdir(parents=True, exist_ok=True)
    variants = read_csv(args.report_dir / "cluster_seen_variant_summary.csv")
    summary = read_csv(args.report_dir / "cluster_seen_summary.csv")

    detailed: dict[tuple[str, str], list[dict[str, object]]] = {}
    for scheme in ("single", "dual"):
        for arm in ("left", "right"):
            selected = [
                row
                for row in variants
                if row["scheme"] == scheme
                and row["arm"] == arm
                and int(row["step"]) == args.step
            ]
            selected.sort(key=lambda row: row["requested_atoms"])
            detailed[(scheme, arm)] = normalized_rows(selected)

    main_rows: dict[str, list[dict[str, object]]] = {}
    for scheme in ("single", "dual"):
        rows = []
        for arm in ("left", "right"):
            match = next(
                row
                for row in summary
                if row["scheme"] == scheme
                and row["arm"] == arm
                and int(row["step"]) == args.step
            )
            rows.append(
                {
                    "canonical": arm.capitalize() + " arm",
                    "symbol": "L" if arm == "left" else "R",
                    "trials": int(match["trials"]),
                    "paired_trials": int(match["paired_trials"]),
                    "reverse": pct(match, "pair_steer_success_rate"),
                    "empty": pct(match, "steer_vs_empty_success_rate"),
                    "absolute": pct(match, "prompt_absolute_success_rate"),
                }
            )
        main_rows[scheme] = rows

    workbook = Workbook()
    workbook.remove(workbook.active)
    add_sheet(workbook, "Main_Single", main_rows["single"])
    add_sheet(workbook, "Main_Dual", main_rows["dual"])
    for scheme in ("single", "dual"):
        for arm in ("left", "right"):
            add_sheet(
                workbook,
                f"{scheme.capitalize()}_{arm.capitalize()}",
                detailed[(scheme, arm)],
            )
    workbook.save(args.output_dir / "atomic_steering_article_tables.xlsx")

    for scheme in ("single", "dual"):
        write_csv(args.output_dir / f"main_{scheme}.csv", main_rows[scheme])
        for arm in ("left", "right"):
            write_csv(
                args.output_dir / f"{scheme}_{arm}_detailed.csv",
                detailed[(scheme, arm)],
            )

    markdown_parts = [
        "# Atomic steering results for publication",
        "",
        "Evaluation contract: target2058 ZM 40K; horizon step 50; translation threshold 5 mm; rotation threshold 1 degree; 250 clusters x 4 anchors; cluster-seen prompts; identical observation, state, and flow noise within each comparison.",
        "",
        "Dual success requires both requested components to pass the threshold simultaneously.",
        "",
        "## Main table A: Single-atom steering",
        "",
        markdown_table(main_rows["single"]),
        "",
        "## Main table B: Dual-atom steering",
        "",
        markdown_table(main_rows["dual"]),
    ]
    for scheme in ("single", "dual"):
        for arm in ("left", "right"):
            markdown_parts.extend(
                [
                    "",
                    f"## Appendix: {scheme.capitalize()} / {arm.capitalize()} arm",
                    "",
                    markdown_table(detailed[(scheme, arm)]),
                ]
            )
    (args.output_dir / "atomic_steering_article_tables.md").write_text(
        "\n".join(markdown_parts) + "\n", encoding="utf-8"
    )

    latex_parts = [
        "% Requires: \\usepackage{booktabs,longtable}",
        latex_table(
            "Single-atom steering success by arm.",
            "tab:atomic-single-main",
            main_rows["single"],
        ),
        latex_table(
            "Dual-atom steering success by arm. Both requested components must pass.",
            "tab:atomic-dual-main",
            main_rows["dual"],
        ),
    ]
    for scheme in ("single", "dual"):
        for arm in ("left", "right"):
            latex_parts.append(
                latex_table(
                    f"{scheme.capitalize()}-atom steering on the {arm} arm.",
                    f"tab:atomic-{scheme}-{arm}",
                    detailed[(scheme, arm)],
                )
            )
    (args.output_dir / "atomic_steering_article_tables.tex").write_text(
        "\n\n".join(latex_parts) + "\n", encoding="utf-8"
    )


if __name__ == "__main__":
    main()
